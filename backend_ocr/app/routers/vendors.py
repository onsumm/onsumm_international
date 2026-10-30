"""Vendor master: listing, creation, edits, PAN de-duplication and merge.

Moved verbatim out of app/main.py; see docs/BACKEND_STRUCTURE.md.
"""
from __future__ import annotations

import psycopg2
from psycopg2.extras import RealDictCursor
from fastapi import FastAPI, File, Form, UploadFile, HTTPException, Query, Request, Depends
from pydantic import BaseModel
from typing import Any, List, Optional
from ..config import DB_URL, TENANT_ID, OPENAI_API_KEY, SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD, SMTP_ADMIN_EMAIL, APPROVAL_SECRET, JWT_SECRET, APP_URL, ACCESS_TOKEN_TTL_MIN, REFRESH_TOKEN_TTL_DAYS, get_conn, put_conn
from ..domain.vendor_merge import repoint_and_retire
from ..security.auth import auth_context, require_tenant, current_tenant, current_role, current_username, require_role
from ..core.vendor_activity_log import log_vendor_activity
from ..domain.vendor_activity import change_title, diff_fields, display_value
import logging as _logging
import uuid as _uuid

from ..core.jurisdiction import jurisdiction_for_cursor
from ..domain.vendor_verification import derive_verification, load_verification_inputs
from fastapi import APIRouter

router = APIRouter()

from ..core.vendor_identity import _effective_pan

@router.get("/api/vendors")
def list_vendors(company_id: str = Query(default=None), brand_id: str = Query(default=None)):
    tenant = require_tenant()  # tenant from JWT; client company_id ignored
    try:
        vendors = _query_vendors(tenant, brand_id)
        return {"vendors": vendors, "total": len(vendors)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/api/vendors/{supplier_id}")
def get_vendor(supplier_id: str, brand_id: str = Query(default=None)):
    """One vendor, in exactly the shape of a /api/vendors row.

    Lets the vendor card open by id from anywhere in the app, without pulling
    the whole list first. Totals are brand-scoped the same way the list is."""
    tenant = require_tenant()
    try:
        _uuid.UUID(supplier_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Vendor not found")
    try:
        vendors = _query_vendors(tenant, brand_id, supplier_id)
        # A card opened by id (from Review, a link) asks for one vendor, not
        # "the vendors of this brand". One that belongs to another of the
        # tenant's brands still exists — show it, with its totals across brands
        # rather than zeros for the brand being viewed.
        if not vendors and brand_id:
            vendors = _query_vendors(tenant, None, supplier_id)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    if not vendors:
        raise HTTPException(status_code=404, detail="Vendor not found")
    return vendors[0]


def _query_vendors(tenant: str, brand_id: str | None, supplier_id: str | None = None) -> list[dict]:
    conn = get_conn()
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute(
            """
            SELECT
                s.supplier_id,
                s.name,
                s.legal_name,
                s.gstn,
                s.email,
                s.payment_terms_days,
                s.status,
                s.is_active,
                s.source,
                s.created_at,
                s.bank_account,
                s.bank_ifsc,
                s.bank_verified,
                s.pan,
                s.pan_dup_deferred,
                -- True when another active vendor in the tenant shares this PAN.
                -- PAN = the GSTIN's embedded PAN (chars 3-12) when a valid
                -- 15-char GSTIN exists (authoritative), else the pan column.
                -- Drives the "Deduplicate" shortcut for same-company records.
                EXISTS (
                    SELECT 1 FROM dim_suppliers s2
                    WHERE s2.tenant_id = s.tenant_id
                      AND s2.is_active = true
                      AND s2.supplier_id <> s.supplier_id
                      AND upper(COALESCE(CASE WHEN length(trim(s.gstn)) = 15 THEN substring(trim(s.gstn) from 3 for 10) END, NULLIF(trim(s.pan), ''))) IS NOT NULL
                      AND upper(COALESCE(CASE WHEN length(trim(s2.gstn)) = 15 THEN substring(trim(s2.gstn) from 3 for 10) END, NULLIF(trim(s2.pan), '')))
                        = upper(COALESCE(CASE WHEN length(trim(s.gstn))  = 15 THEN substring(trim(s.gstn)  from 3 for 10) END, NULLIF(trim(s.pan),  '')))
                ) AS has_pan_duplicate,
                s.vendor_category_id,
                vc.name                          AS vendor_category_name,
                s.brand_id,
                COUNT(fi.invoice_id)          AS total_bills,
                COALESCE(SUM(fi.total_amount), 0) AS total_billed,
                COALESCE(SUM(CASE WHEN fi.payment_status = 'paid' THEN 0 ELSE fi.outstanding_amount END), 0) AS outstanding_amount,
                MAX(fi.invoice_date)          AS last_invoice_date
            FROM dim_suppliers s
            LEFT JOIN fact_invoices fi
                   ON fi.supplier_id = s.supplier_id
                  AND fi.tenant_id   = s.tenant_id
                  -- Exclude deleted / on-hold bills so archived bills never
                  -- inflate the vendor totals. Kept in the ON clause (not the
                  -- WHERE) so vendors with no active bills still appear.
                  AND COALESCE(fi.soft_status, 'active') = 'active'
                  -- Aggregate only the bills of the brand being viewed, so the
                  -- totals are scoped per brand even for a vendor shared across
                  -- brands (and migrating a vendor never carries its bills along).
                  AND (%s::uuid IS NULL OR fi.brand_id = %s::uuid)
            LEFT JOIN dim_vendor_categories vc
                   ON vc.category_id = s.vendor_category_id
            WHERE s.tenant_id = %s
              AND s.is_active = true
              AND (%s::uuid IS NULL OR s.brand_id = %s::uuid OR s.brand_id IS NULL)
              AND (%s::uuid IS NULL OR s.supplier_id = %s::uuid)
            GROUP BY s.supplier_id, s.name, s.legal_name, s.gstn, s.email,
                     s.payment_terms_days, s.status, s.is_active,
                     s.source, s.created_at, s.bank_account, s.bank_ifsc, s.bank_verified, s.pan,
                     s.pan_dup_deferred, s.vendor_category_id, vc.name, s.brand_id
            ORDER BY total_bills DESC, s.name ASC
            """,
            (brand_id, brand_id, tenant, brand_id, brand_id, supplier_id, supplier_id),
        )
        rows = cur.fetchall()
        # The verification badge depends on which country's rules apply
        # (domain/vendor_verification.py): GSTIN + IFSC bank for India, a
        # payable wallet/bank plus tax documents elsewhere.
        jur = jurisdiction_for_cursor(cur, tenant)
        ver_inputs = load_verification_inputs(cur, tenant, [r["supplier_id"] for r in rows])
    put_conn(conn)

    vendors = []
    for r in rows:
        vi = ver_inputs.get(str(r["supplier_id"]), {})
        vendors.append({
            "supplier_id":       str(r["supplier_id"]),
            "name":              r["name"],
            # Brand name is `name`; the registered (GST) name is legal_name, the
            # same as name until the GST portal says otherwise.
            "legal_name":        r["legal_name"] or r["name"],
            "gstn":              r["gstn"].strip() if r["gstn"] else None,
            "email":             r["email"],
            "payment_terms_days": r["payment_terms_days"],
            "status":            r["status"],
            "is_active":         r["is_active"],
            "source":            r["source"],
            "total_bills":       r["total_bills"],
            "total_billed":      float(r["total_billed"]),
            "outstanding_amount": float(r["outstanding_amount"]),
            "last_invoice_date": r["last_invoice_date"].isoformat() if r["last_invoice_date"] else None,
            "created_at":        r["created_at"].isoformat() if r["created_at"] else None,
            "bank_account":      r["bank_account"],
            "bank_ifsc":         r["bank_ifsc"],
            "bank_verified":     r["bank_verified"] or False,
            "pan":               r["pan"] if r["pan"] else (r["gstn"][2:12] if r["gstn"] and len(r["gstn"]) == 15 else None),
            "pan_dup_deferred":  bool(r["pan_dup_deferred"]),
            "has_pan_duplicate": bool(r["has_pan_duplicate"]),
            "vendor_category_id":   str(r["vendor_category_id"]) if r["vendor_category_id"] else None,
            "vendor_category_name": r["vendor_category_name"],
            "brand_id":             str(r["brand_id"]) if r["brand_id"] else None,
            # Documents the vendor has given us (W-8BEN-E, TRC, tax ID...) and where
            # they are resident for tax: what the Compliance tab lists for a
            # tenant without GST, and what "verified" depends on there.
            "tax_docs":             vi.get("tax_docs") or {},
            "country_code":         vi.get("country_code"),
            "tax_residency":        vi.get("tax_residency"),
            "verification": derive_verification(
                jur,
                gst_status=r["status"],
                bank_verified=bool(r["bank_verified"]) or vi.get("has_verified_bank_row", False),
                has_active_wallet=vi.get("has_active_wallet", False),
                tax_docs=vi.get("tax_docs"),
            ),
        })
    return vendors

class VendorCreateRequest(BaseModel):
    name: str
    gstn: str | None = None
    email: str | None = None
    payment_terms_days: int = 30
    company_id: str | None = None
    bank_account: str | None = None
    bank_ifsc: str | None = None
    bank_name: str | None = None
    # ISO 3166 alpha-2. A tenant without the GST pack identifies a vendor by
    # country instead of GSTIN; an Indian tenant's vendors are IN.
    country_code: str | None = None

@router.post("/api/vendors")
def create_vendor(body: VendorCreateRequest):
    tenant = require_tenant()  # tenant from JWT; client company_id ignored
    import re as _re
    GSTN_RE = r'^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]{1}[1-9A-Z]{1}Z[0-9A-Z]{1}$'
    acct = (body.bank_account or "").strip()
    ifsc = (body.bank_ifsc or "").strip().upper()
    name = body.name.strip()
    if not name:
        raise HTTPException(status_code=422, detail="Vendor name is required")
    gstn_clean = (body.gstn or "").strip().upper() or None
    country = (body.country_code or "").strip().upper() or None
    if country and not _re.match(r'^[A-Z]{2}$', country):
        raise HTTPException(status_code=422, detail="Country must be a 2-letter ISO code (e.g. US, SG)")

    # Vendor is saved as UNVERIFIED; the verification pipeline handles KYB asynchronously.
    conn = get_conn()
    conn.autocommit = False
    try:
        with conn.cursor() as cur:
            # GSTIN is the identity of an Indian vendor and is required there.
            # Elsewhere there is none to ask for: it is optional, and the
            # vendor is placed by country.
            if jurisdiction_for_cursor(cur, tenant).has("gst"):
                if not gstn_clean:
                    raise HTTPException(status_code=422, detail="GSTN is required")
                if not _re.match(GSTN_RE, gstn_clean):
                    raise HTTPException(status_code=422, detail="Invalid GSTN format (e.g. 29AABCT1332L1ZU)")
                country = country or "IN"
            elif gstn_clean and not _re.match(GSTN_RE, gstn_clean):
                raise HTTPException(status_code=422, detail="Invalid GSTN format (e.g. 29AABCT1332L1ZU)")
            source_id = gstn_clean or f"MANUAL_{name.lower().replace(' ', '_')}_{str(_uuid.uuid4())[:8]}"
            cur.execute(
                """INSERT INTO dim_suppliers
                   (tenant_id, name, gstn, email, payment_terms_days, status, is_active, source, source_supplier_id,
                    bank_account, bank_ifsc, bank_verified, bank_name, country_code, tax_residency)
                   VALUES (%s, %s, %s, %s, %s, 'UNVERIFIED', true, 'MANUAL', %s, %s, %s, false, %s, %s, %s)
                   RETURNING supplier_id""",
                (tenant, name, gstn_clean, body.email, body.payment_terms_days, source_id,
                 acct or None, ifsc or None, body.bank_name or None, country, country),
            )
            supplier_id = str(cur.fetchone()[0])
        conn.commit()
        return {
            "supplier_id": supplier_id,
            "status": "created",
            "verification_status": "UNVERIFIED",
        }
    except HTTPException:
        conn.rollback()
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        put_conn(conn)

class VendorMigrateRequest(BaseModel):
    company_id: str
    brand_id: str | None = None  # None = shared across all brands
    supplier_ids: list[str] | None = None  # for bulk migration

@router.patch("/api/vendors/{supplier_id}/brand")
def migrate_vendor_brand(supplier_id: str, body: VendorMigrateRequest):
    """Move one or more vendors to a specific brand or make them shared (brand_id=null)."""
    tenant = require_tenant()  # tenant from JWT; client company_id ignored
    ids = body.supplier_ids if body.supplier_ids else [supplier_id]
    conn = get_conn()
    try:
        with conn.cursor() as cur:
            for sid in ids:
                cur.execute(
                    "UPDATE dim_suppliers SET brand_id = %s::uuid WHERE supplier_id = %s::uuid AND tenant_id = %s",
                    (body.brand_id, sid, tenant),
                )
        conn.commit()
        return {"status": "ok", "brand_id": body.brand_id, "updated": len(ids)}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        put_conn(conn)

class VendorUpdateRequest(BaseModel):
    name: Optional[str] = None
    legal_name: Optional[str] = None
    email: Optional[str] = None
    payment_terms_days: Optional[int] = None
    pan: Optional[str] = None
    vendor_category_id: Optional[str] = None

class VendorNamesRequest(BaseModel):
    name: str
    legal_name: str


@router.post("/api/vendors/{supplier_id}/names")
def set_vendor_names(supplier_id: str, body: VendorNamesRequest):
    """Set a vendor's brand name and legal name together.

    The Review action for a vendor whose name the GST portal replaced: keep the
    portal's name as the legal name and restore the name the business uses as
    the brand name. One write, so the two never disagree half-way.
    """
    tenant = require_tenant()
    name, legal = body.name.strip(), body.legal_name.strip()
    if not name or not legal:
        raise HTTPException(status_code=400, detail="Both names are required")
    conn = get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT name, legal_name FROM dim_suppliers"
                " WHERE supplier_id = %s AND tenant_id = %s FOR UPDATE",
                (supplier_id, tenant),
            )
            before = cur.fetchone()
            cur.execute(
                "UPDATE dim_suppliers SET name = %s, legal_name = %s"
                " WHERE supplier_id = %s AND tenant_id = %s",
                (name, legal, supplier_id, tenant),
            )
            if cur.rowcount == 0:
                raise HTTPException(status_code=404, detail="Vendor not found")
            _log_vendor_changes(cur, tenant, supplier_id,
                                {"name": before[0], "legal_name": before[1]} if before else {},
                                {"name": name, "legal_name": legal})
            # Re-check now rather than on the next unrelated write: the finding
            # this answers clears as soon as the name is no longer the portal's.
            cur.execute("SELECT pg_notify('ch_review_adapters', '')")
        conn.commit()
    except HTTPException:
        conn.rollback()
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        put_conn(conn)
    return {"ok": True, "name": name, "legal_name": legal}


def _log_vendor_changes(cur, tenant: str, supplier_id: str, before: dict, after: dict) -> None:
    """Write a 'vendor_updated' Activity row with what actually changed (if anything)."""
    category_names: dict = {}
    if "vendor_category_id" in after:
        ids = [str(v) for v in (before.get("vendor_category_id"), after.get("vendor_category_id")) if v]
        if ids:
            try:
                cur.execute("SAVEPOINT vendor_cat_names")
                cur.execute("SELECT category_id::text, name FROM dim_vendor_categories"
                            " WHERE category_id::text = ANY(%s)", (ids,))
                category_names = {r[0]: r[1] for r in cur.fetchall()}
                cur.execute("RELEASE SAVEPOINT vendor_cat_names")
            except Exception:
                cur.execute("ROLLBACK TO SAVEPOINT vendor_cat_names")

    def fmt(field, value):
        if field == "vendor_category_id" and value:
            return category_names.get(str(value), "another category")
        return display_value(field, value)

    changes = diff_fields(before, after, value_fmt=fmt)
    if changes:
        log_vendor_activity(cur, supplier_id, "vendor_updated", change_title(changes),
                            changes=changes, tenant_id=tenant)


@router.patch("/api/vendors/{supplier_id}")
def update_vendor(supplier_id: str, body: VendorUpdateRequest, company_id: Optional[str] = None):
    tenant = require_tenant()  # tenant from JWT; client company_id ignored
    # Verify the vendor belongs to the caller's tenant (no cross-tenant edits).
    try:
        _rc = get_conn()
        with _rc.cursor() as _cur:
            _cur.execute("SELECT tenant_id FROM dim_suppliers WHERE supplier_id = %s", (supplier_id,))
            _row = _cur.fetchone()
        put_conn(_rc)
        if _row and str(_row[0]) != tenant:
            raise HTTPException(status_code=404, detail="Vendor not found")
    except HTTPException:
        raise
    except Exception:
        pass
    fields: dict = {}
    if body.name is not None:
        fields["name"] = body.name.strip()
    if body.legal_name is not None:
        fields["legal_name"] = body.legal_name.strip() or None
    if body.email is not None:
        fields["email"] = body.email.strip() or None
    if body.payment_terms_days is not None:
        fields["payment_terms_days"] = body.payment_terms_days
    if body.pan is not None:
        fields["pan"] = body.pan.strip().upper() or None
    category_changed = body.vendor_category_id is not None
    if category_changed:
        fields["vendor_category_id"] = body.vendor_category_id or None
    if not fields:
        raise HTTPException(status_code=400, detail="No fields to update")
    set_clause = ", ".join(f"{k} = %s" for k in fields)
    conn = psycopg2.connect(DB_URL)
    try:
        with conn.cursor() as cur:
            # The row as it was, for the Activity log's before → after. Column
            # names come from the fixed keys above, never from the request.
            cur.execute(
                f"SELECT {', '.join(fields)} FROM dim_suppliers "
                "WHERE supplier_id = %s AND tenant_id = %s FOR UPDATE",
                (supplier_id, tenant),
            )
            before_row = cur.fetchone()
            cur.execute(
                f"UPDATE dim_suppliers SET {set_clause} "
                "WHERE supplier_id = %s AND tenant_id = %s",
                (*fields.values(), supplier_id, tenant),
            )
            if cur.rowcount == 0:
                raise HTTPException(status_code=404, detail="Vendor not found")
            _log_vendor_changes(cur, tenant, supplier_id,
                                dict(zip(fields, before_row)) if before_row else {}, fields)
            if category_changed:
                # These bills' amounts are already baked into vendor_tds_fy_state's
                # cumulative totals. calculate_tds() always adds on top of the stored
                # cumulative rather than replacing it, so reprocessing them without
                # clearing that state first would double-count every bill. Since every
                # NA/CALCULATED bill for this vendor is about to be requeued, wiping all
                # of its FY state rows is a full, correct reset (not a partial one).
                cur.execute(
                    "DELETE FROM vendor_tds_fy_state WHERE tenant_id=%s AND supplier_id=%s",
                    (tenant, supplier_id),
                )
                cur.execute(
                    "UPDATE stg_bills SET tds_status='PENDING'"
                    " WHERE supplier_id=%s AND tenant_id=%s"
                    " AND tds_status IN ('NA','CALCULATED')",
                    (supplier_id, tenant),
                )
                reset_count = cur.rowcount
            else:
                reset_count = 0
        conn.commit()
        return {"ok": True, "bills_reset": reset_count}
    except HTTPException:
        raise
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()

@router.get("/api/vendors/{supplier_id}/pan-duplicates")
def vendor_pan_duplicates(supplier_id: str, company_id: Optional[str] = None):
    """Other active vendors in the same tenant that share this vendor's PAN.
    A company has one PAN but may hold several GSTINs, so PAN is the identity
    key used to detect duplicate vendor records (e.g. one created from an
    OCR-scrambled GSTIN)."""
    tenant = require_tenant()
    conn = get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT tenant_id, name, gstn, pan FROM dim_suppliers WHERE supplier_id=%s AND tenant_id=%s",
                (supplier_id, tenant),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="Vendor not found")
            tenant_id, _name, gstn, pan = row
            eff_pan = _effective_pan(gstn, pan)
            if not eff_pan:
                return {"pan": None, "duplicates": []}
            cur.execute(
                """
                SELECT s.supplier_id, s.name, s.gstn, s.pan, s.status,
                       (SELECT count(*) FROM fact_invoices fi WHERE fi.supplier_id = s.supplier_id) AS bill_count
                FROM   dim_suppliers s
                WHERE  s.tenant_id = %s
                  AND  s.is_active = true
                  AND  s.supplier_id <> %s
                  AND  upper(COALESCE(CASE WHEN length(trim(s.gstn)) = 15 THEN substring(trim(s.gstn) from 3 for 10) END, NULLIF(trim(s.pan), ''))) = %s
                ORDER  BY bill_count DESC, s.created_at ASC
                """,
                (tenant_id, supplier_id, eff_pan),
            )
            dups = [
                {
                    "supplier_id": str(r[0]),
                    "name": r[1],
                    "gstn": r[2],
                    "pan": r[3] or (r[2][2:12].upper() if r[2] and len(r[2]) == 15 else None),
                    "status": r[4],
                    "bill_count": r[5],
                }
                for r in cur.fetchall()
            ]
            return {"pan": eff_pan, "duplicates": dups}
    finally:
        put_conn(conn)

class VendorMergeRequest(BaseModel):
    target_supplier_id: str
    actor_name: str | None = None
    actor_role: str | None = None

@router.post("/api/vendors/{supplier_id}/merge")
def merge_vendor(supplier_id: str, body: VendorMergeRequest):
    """Merge the vendor `supplier_id` (source / duplicate) INTO `target_supplier_id`
    (survivor). All rows referencing the source's supplier_id are repointed to the
    target; rows that would collide with an existing target row (e.g. a duplicate
    GSTIN in supplier_gstns) are dropped instead. The source vendor is then
    deactivated (is_active=false) — kept for audit, hidden from the vendor list.

    Guardrail: both vendors must be in the same tenant and share the same PAN."""
    if body.target_supplier_id == supplier_id:
        raise HTTPException(status_code=400, detail="Cannot merge a vendor into itself")

    tenant = require_tenant()

    conn = get_conn()
    conn.autocommit = False
    try:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT tenant_id, name, gstn, pan, is_active FROM dim_suppliers WHERE supplier_id=%s AND tenant_id=%s",
                (supplier_id, tenant),
            )
            src = cur.fetchone()
            cur.execute(
                "SELECT tenant_id, name, gstn, pan FROM dim_suppliers WHERE supplier_id=%s AND tenant_id=%s",
                (body.target_supplier_id, tenant),
            )
            tgt = cur.fetchone()
            if not src or not tgt:
                raise HTTPException(status_code=404, detail="Vendor not found")
            # Both rows already matched the caller's tenant above, so this can no
            # longer differ — kept as a belt-and-braces assertion on the merge.
            if str(src[0]) != str(tgt[0]):
                raise HTTPException(status_code=400, detail="Vendors belong to different tenants")

            src_pan = _effective_pan(src[2], src[3])
            tgt_pan = _effective_pan(tgt[2], tgt[3])
            if not src_pan or not tgt_pan or src_pan != tgt_pan:
                raise HTTPException(
                    status_code=400,
                    detail=f"PAN mismatch — refusing to merge (source={src_pan}, target={tgt_pan})",
                )

            # Shared with the Review Hub's misread-GSTIN fix — see
            # app/domain/vendor_merge.py for why the guardrails stay out here.
            moved = repoint_and_retire(
                cur, supplier_id, body.target_supplier_id,
                source_gstn=src[2], source_pan=src_pan,
            )
        conn.commit()
        _logging.getLogger(__name__).info(
            "[vendor_merge] %s (%s) merged into %s (%s) by %s; moved=%s",
            supplier_id, src[1], body.target_supplier_id, tgt[1], body.actor_name, moved,
        )
        return {
            "status": "merged",
            "source_supplier_id": supplier_id,
            "target_supplier_id": body.target_supplier_id,
            "target_name": tgt[1],
            "moved": moved,
        }
    except HTTPException:
        conn.rollback(); raise
    except Exception as e:
        conn.rollback(); raise HTTPException(status_code=500, detail=str(e))
    finally:
        put_conn(conn)

class DeferDedupRequest(BaseModel):
    deferred: bool = True

@router.post("/api/vendors/{supplier_id}/defer-dedup")
def defer_vendor_dedup(supplier_id: str, body: DeferDedupRequest):
    """Record that the user chose to KEEP a PAN-duplicate vendor separate (or clear
    that choice). When set, the vendor list shows a green 'Deduplicate' shortcut."""
    tenant = require_tenant()
    conn = get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE dim_suppliers SET pan_dup_deferred=%s WHERE supplier_id=%s AND tenant_id=%s",
                (body.deferred, supplier_id, tenant),
            )
            if cur.rowcount == 0:
                raise HTTPException(status_code=404, detail="Vendor not found")
        conn.commit()
        return {"status": "ok", "supplier_id": supplier_id, "pan_dup_deferred": body.deferred}
    except HTTPException:
        conn.rollback(); raise
    except Exception as e:
        conn.rollback(); raise HTTPException(status_code=500, detail=str(e))
    finally:
        put_conn(conn)

class OutreachRequest(BaseModel):
    requested_by: str
    email_to: str | None = None
    company_id: str | None = None

@router.post("/api/vendors/{supplier_id}/request-info")
def request_vendor_info(supplier_id: str, body: OutreachRequest):
    """Send a missing-info request email to the vendor and log it."""
    tenant = require_tenant()  # tenant from JWT; client company_id ignored

    conn = get_conn()
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT name, gstn, email FROM dim_suppliers WHERE supplier_id = %s AND tenant_id = %s",
                        (supplier_id, tenant))
            vendor = cur.fetchone()
    finally:
        put_conn(conn)

    if not vendor:
        raise HTTPException(status_code=404, detail="Vendor not found")

    email_to = body.email_to or vendor["email"]
    subject = f"Action Required: Please update your compliance details — {vendor['name']}"
    body_text = (
        f"Dear {vendor['name']},\n\n"
        "We are reaching out to collect/verify your compliance details required for payment processing.\n\n"
        "Please provide the following:\n"
        f"{'- GSTIN (GST Identification Number)' if not vendor['gstn'] else ''}\n"
        "- PAN (Permanent Account Number)\n"
        "- Bank Account Number and IFSC Code\n\n"
        "Kindly reply to this email with the required details at your earliest convenience.\n\n"
        "Best regards,\nAccounts Team"
    )

    conn = get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO vendor_outreach_log
                   (supplier_id, tenant_id, requested_by, email_to, subject, body, status)
                   VALUES (%s, %s, %s, %s, %s, %s, 'logged')""",
                (supplier_id, tenant, body.requested_by, email_to, subject, body_text),
            )
        conn.commit()
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        put_conn(conn)

    return {"status": "logged", "email_to": email_to, "subject": subject, "body": body_text}

@router.get("/api/vendors/{supplier_id}/outreach")
def get_vendor_outreach(supplier_id: str, company_id: str = Query(default=None)):
    """Get outreach history for a vendor."""
    tenant = require_tenant()  # tenant from JWT; client company_id ignored

    conn = get_conn()
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(
                "SELECT id, requested_by, requested_at, email_to, subject, status FROM vendor_outreach_log "
                "WHERE supplier_id = %s AND tenant_id = %s ORDER BY requested_at DESC",
                (supplier_id, tenant),
            )
            rows = cur.fetchall()
    finally:
        put_conn(conn)

    return {"logs": [
        {**dict(r), "id": str(r["id"]), "requested_at": r["requested_at"].isoformat() if r["requested_at"] else None}
        for r in rows
    ]}

@router.post("/api/vendors/bulk-excel")
async def bulk_create_vendors(file: UploadFile = File(...), company_id: str = Query(default=None)):
    import psycopg2, io
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    tenant = require_tenant()  # tenant from JWT; client company_id ignored
    content = await file.read()
    fname = (file.filename or '').lower()

    if fname.endswith('.csv'):
        import csv as csv_mod
        text = content.decode('utf-8-sig')
        reader = list(csv_mod.DictReader(io.StringIO(text)))
        raw_headers = list(reader[0].keys()) if reader else []
        headers = [h.strip().lower().replace(' ', '_') for h in raw_headers]
        class _Row:
            def __init__(self, d): self._d = d
            def __getitem__(self, i): return type('C', (), {'value': list(self._d.values())[i]})()
        rows_data = [_Row(r) for r in reader]
    else:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        headers = [str(c.value).strip().lower().replace(' ', '_') if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
        rows_data = list(ws.iter_rows(min_row=2, values_only=False))

    def col(row, name):
        alts = {'vendor_name': 'name', 'gstin': 'gstn', 'payment_terms': 'payment_terms_days'}
        key = alts.get(name, name)
        for h in [name, key]:
            if h in headers:
                v = row[headers.index(h)].value
                return str(v).strip() if v is not None else None
        return None

    created, skipped, errors = 0, 0, []
    conn = get_conn()
    conn.autocommit = False
    try:
        with conn.cursor() as cur:
            import uuid as _uuid
            for i, row in enumerate(rows_data, start=2):
                name = col(row, 'name') or col(row, 'vendor_name')
                if not name:
                    skipped += 1
                    continue
                gstn_raw = col(row, 'gstn') or col(row, 'gstin')
                gstn = gstn_raw.strip().upper() if gstn_raw else None
                email = col(row, 'email')
                try:
                    terms = int(col(row, 'payment_terms_days') or col(row, 'payment_terms') or 30)
                except (ValueError, TypeError):
                    terms = 30
                source_id = gstn if gstn else f"MANUAL_{name.strip().lower().replace(' ', '_')}_{str(_uuid.uuid4())[:8]}"
                try:
                    cur.execute(
                        """INSERT INTO dim_suppliers
                           (tenant_id, name, gstn, email, payment_terms_days, status, is_active, source, source_supplier_id)
                           VALUES (%s,%s,%s,%s,%s,'UNVERIFIED',true,'MANUAL',%s)
                           ON CONFLICT DO NOTHING""",
                        (tenant, name, gstn, email, terms, source_id),
                    )
                    if cur.rowcount:
                        created += 1
                    else:
                        skipped += 1
                except Exception as ex:
                    errors.append(f"Row {i}: {ex}")
        conn.commit()
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        put_conn(conn)

    return {"created": created, "skipped": skipped, "errors": errors}

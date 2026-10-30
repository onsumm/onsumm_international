"""Expense reporting and its spreadsheet/CSV exports.

Moved verbatim out of app/main.py; see docs/BACKEND_STRUCTURE.md.
"""
from __future__ import annotations

import psycopg2
from psycopg2.extras import RealDictCursor
from fastapi import FastAPI, File, Form, UploadFile, HTTPException, Query, Request, Depends
from ..config import DB_URL, TENANT_ID, OPENAI_API_KEY, SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD, SMTP_ADMIN_EMAIL, APPROVAL_SECRET, JWT_SECRET, APP_URL, ACCESS_TOKEN_TTL_MIN, REFRESH_TOKEN_TTL_DAYS, get_conn, put_conn
from ..core.jurisdiction import get_jurisdiction
from ..security.auth import auth_context, require_tenant, current_tenant, current_role, current_username, require_role

from fastapi import APIRouter

router = APIRouter()

@router.get("/api/reports/activity")
def activity_report(company_id: str | None = Query(None), brand_id: str | None = Query(None)):
    """
    Per-bill lifecycle, extracted from the activity trail.

    Everything Reports needs about timing already exists in `bill_approval_chain`
    — the same rows the per-bill "View activity" timeline renders. Until now it
    was only reachable one bill at a time, so a book-wide turnaround chart would
    have meant an N+1 across the whole ledger. This collapses it to one query and
    returns one compact row per bill:

      first_at    when the bill entered the chain (first recorded action)
      approved_at when it cleared its LAST approval — the timestamp that ends
                  turnaround and starts the payment buffer
      send_backs  how many times it was pushed back for revision, i.e. rework
      role_days   days the bill sat with each desk before that desk acted, so
                  "where does the time go" is answerable per stage rather than
                  as one opaque total

    A stage's wait is measured from the previous action on that bill, falling back
    to the bill's own created_at for the first action — otherwise the first
    approver's queue time would be invisible. Clamped at zero: a clock-skewed or
    backdated chain row is missing data, not a negative wait.

    Step 99 is excluded throughout; it is the reserved payment/audit step, not a
    stage anybody waits in.
    """

    tenant = require_tenant()  # tenant from JWT; client company_id ignored

    conn = get_conn()
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(
                """
                WITH ev AS (
                    SELECT
                        c.invoice_id,
                        c.action,
                        c.actor_role,
                        c.created_at,
                        COALESCE(
                            LAG(c.created_at) OVER (
                                PARTITION BY c.invoice_id ORDER BY c.created_at, c.id
                            ),
                            fi.created_at
                        ) AS prev_at
                    FROM bill_approval_chain c
                    JOIN fact_invoices fi
                          ON fi.invoice_id = c.invoice_id
                         AND fi.tenant_id  = c.tenant_id
                    WHERE c.tenant_id = %s
                      AND c.step <> 99
                      AND (%s::uuid IS NULL OR fi.brand_id = %s::uuid)
                ),
                per_role AS (
                    SELECT
                        invoice_id,
                        COALESCE(NULLIF(TRIM(actor_role), ''), 'Unattributed') AS role,
                        SUM(
                            GREATEST(EXTRACT(EPOCH FROM (created_at - prev_at)), 0)
                        ) / 86400.0 AS days
                    FROM ev
                    GROUP BY 1, 2
                ),
                role_json AS (
                    SELECT invoice_id,
                           jsonb_object_agg(role, ROUND(days::numeric, 4)) AS role_days
                    FROM per_role
                    GROUP BY 1
                )
                SELECT
                    ev.invoice_id::text                        AS invoice_id,
                    MIN(ev.created_at)                         AS first_at,
                    MAX(ev.created_at) FILTER (
                        WHERE ev.action IN ('approve', 'approved')
                    )                                          AS approved_at,
                    COUNT(*) FILTER (
                        WHERE ev.action IN ('send_back', 'sent_back', 'reject')
                    )::int                                     AS send_backs,
                    COALESCE(rj.role_days, '{}'::jsonb)        AS role_days
                FROM ev
                LEFT JOIN role_json rj ON rj.invoice_id = ev.invoice_id
                GROUP BY ev.invoice_id, rj.role_days
                """,
                (tenant, brand_id, brand_id),
            )
            rows = cur.fetchall()
    finally:
        put_conn(conn)

    def iso(v):
        return v.isoformat() if v is not None else None

    return {
        "rows": [
            {
                "invoice_id": r["invoice_id"],
                "first_at": iso(r["first_at"]),
                "approved_at": iso(r["approved_at"]),
                "send_backs": r["send_backs"],
                "role_days": r["role_days"] or {},
            }
            for r in rows
        ],
    }

@router.get("/api/reports/expense-xlsx")
def expense_xlsx(
    company_id: str | None = Query(None),
    month: str | None = Query(None),   # e.g. "2026-04"
    status: str | None = Query(None),
    brand_id: str | None = Query(None),
):
    import psycopg2.extras
    import io
    from datetime import date
    from calendar import monthrange
    from fastapi.responses import StreamingResponse
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    tenant = require_tenant()  # tenant from JWT; client company_id ignored
    # The books currency (core/jurisdiction.py): INR for Indian tenants. A bill
    # is foreign only when its currency differs from it. The INR-per-unit rates
    # in dim_fx_rates revalue against INR, so the FX columns and the
    # revaluation sheet are produced only for an INR-books tenant.
    base_ccy = get_jurisdiction(tenant).base_currency

    # ── Date range (month=None or "all" → no date filter) ───────────────────
    use_date_filter = bool(month and month.lower() != "all")
    if use_date_filter:
        year, mon = int(month.split("-")[0]), int(month.split("-")[1])
        first_day = date(year, mon, 1)
        next_first = date(year, mon + 1, 1) if mon < 12 else date(year + 1, 1, 1)
        month_label = first_day.strftime("%b-%y").upper()
    else:
        month_label = "ALL"

    conn = get_conn()
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            # Fetch bills
            sql = """
                SELECT
                    fi.invoice_id,
                    fi.stg_bill_id,
                    fi.invoice_date,
                    fi.due_date,
                    fi.total_amount,
                    fi.category,
                    fi.status,
                    fi.payment_status,
                    fi.outstanding_amount,
                    ds.name AS vendor_name,
                    fi.currency_code,
                    fi.original_amount,
                    fi.inr_amount_current,
                    fi.fx_gain_loss,
                    dr.inr_per_unit AS fx_rate,
                    dr_current.inr_per_unit AS fx_rate_current
                FROM fact_invoices fi
                LEFT JOIN dim_suppliers ds ON ds.supplier_id = fi.supplier_id
                LEFT JOIN LATERAL (
                    SELECT inr_per_unit FROM dim_fx_rates
                    WHERE currency_code = fi.currency_code
                      AND rate_date <= COALESCE(fi.invoice_date, CURRENT_DATE)
                    ORDER BY rate_date DESC LIMIT 1
                ) dr ON fi.currency_code IS NOT NULL AND fi.currency_code != 'INR'
                LEFT JOIN LATERAL (
                    SELECT inr_per_unit FROM dim_fx_rates
                    WHERE currency_code = fi.currency_code
                    ORDER BY rate_date DESC LIMIT 1
                ) dr_current ON fi.currency_code IS NOT NULL AND fi.currency_code != 'INR'
                WHERE fi.tenant_id = %s
            """
            params: list = [tenant]
            if use_date_filter:
                sql += " AND fi.invoice_date >= %s AND fi.invoice_date < %s"
                params.extend([first_day, next_first])
            if status:
                sql += " AND UPPER(fi.status) = %s"
                params.append(status.upper())
            if brand_id:
                sql += " AND fi.brand_id = %s::uuid"
                params.append(brand_id)
            cur.execute(sql, params)
            bills = cur.fetchall()

            # Fetch COA rules ordered by priority desc
            cur.execute(
                """
                SELECT exp_head, vendor_name, category, priority
                FROM coa_rules
                WHERE tenant_id = %s
                ORDER BY priority DESC
                """,
                (tenant,)
            )
            rules = cur.fetchall()

            # Fetch dim_categories exp_head assignments (4th tier fallback)
            cur.execute(
                "SELECT name, exp_head FROM dim_categories WHERE tenant_id = %s AND exp_head IS NOT NULL",
                (tenant,)
            )
            cat_exp_heads = {r["name"].lower(): r["exp_head"] for r in cur.fetchall()}
    finally:
        put_conn(conn)

    def resolve_exp_head(vendor: str | None, category: str | None) -> str:
        v = (vendor or "").lower()
        c = (category or "").lower()
        for r in rules:
            rv = (r["vendor_name"] or "").lower()
            rc = (r["category"] or "").lower()
            if rv and rc and rv in v and rc == c:
                return r["exp_head"]
        for r in rules:
            rv = (r["vendor_name"] or "").lower()
            rc = r["category"]
            if rv and not rc and rv in v:
                return r["exp_head"]
        for r in rules:
            rv = r["vendor_name"]
            rc = (r["category"] or "").lower()
            if not rv and rc and rc == c:
                return r["exp_head"]
        # 4th tier: dim_categories exp_head assignment
        if c in cat_exp_heads:
            return cat_exp_heads[c]
        return "OTHER"

    # ── Build workbook ──────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    ws_raw = wb.active
    ws_raw.title = "RAW DATA"

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D3D3D3")

    has_fx_bills = base_ccy == "INR" and any(
        (b.get("currency_code") or "INR") != "INR" and b.get("original_amount") is not None
        for b in bills
    )

    raw_headers = ["COMPANY", "VENDOR NAME", "INVOICE NO", "SERVICE MONTH",
                   "SUB HEAD", "SERVICE", "EXP HEAD", "TOTAL", "OPEX", "MONTHLY EXPENSE"]
    if has_fx_bills:
        raw_headers.extend(["ORIGINAL CURRENCY", "ORIGINAL AMOUNT", "FX RATE"])
    for col_idx, h in enumerate(raw_headers, 1):
        cell = ws_raw.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    company_name = "COMPANY"
    # exp_head → {category → total}
    exp_head_cat_totals: dict[str, dict[str, float]] = {}

    for row_idx, bill in enumerate(bills, 2):
        vendor   = bill["vendor_name"] or ""
        category = (bill["category"] or "General").strip()
        exp_head = resolve_exp_head(vendor, category)
        total    = float(bill["total_amount"] or 0)
        opex     = "Y" if (bill["payment_status"] or "").lower() in ("paid", "partial") else None
        inv_date  = bill["invoice_date"]
        svc_month = inv_date.strftime("%b-%y").upper() if inv_date else month_label
        inv_no    = str(bill["stg_bill_id"] or bill["invoice_id"] or "")

        row_data = [company_name, vendor, inv_no, svc_month,
                    vendor, category, exp_head, total, opex, total]
        if has_fx_bills:
            ccy = (bill.get("currency_code") or "INR").strip()
            orig = float(bill["original_amount"]) if bill.get("original_amount") is not None else None
            rate = float(bill["fx_rate"]) if bill.get("fx_rate") is not None else None
            row_data.extend([ccy, orig, rate])
        ws_raw.append(row_data)

        eh_cats = exp_head_cat_totals.setdefault(exp_head, {})
        eh_cats[category] = eh_cats.get(category, 0.0) + total

    # ── SUMMARY sheet ───────────────────────────────────────────────────────
    ws_sum = wb.create_sheet("SUMMARY")
    sum_headers = ["EXP HEAD", "CATEGORY", "AMOUNT", "EXP RATIO IN %"]
    for col_idx, h in enumerate(sum_headers, 1):
        cell = ws_sum.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    grand_total = sum(sum(cats.values()) for cats in exp_head_cat_totals.values()) or 1

    # Style helpers
    from openpyxl.styles import Border, Side
    _hf = PatternFill("solid", fgColor="1E3A8A")   # dark indigo — exp head header
    _hft = Font(bold=True, color="FFFFFF")
    _tf  = PatternFill("solid", fgColor="DBEAFE")  # light blue  — total row
    _tft = Font(bold=True, color="1E3A8A")
    _cf  = PatternFill("solid", fgColor="F0F4FF")  # very light  — category rows
    _gtf = PatternFill("solid", fgColor="1E293B")  # slate dark  — grand total
    _gtft= Font(bold=True, color="FFFFFF")

    sorted_heads = sorted(exp_head_cat_totals.keys(), key=lambda h: (h == "OTHER", h))
    ri = 2
    for eh in sorted_heads:
        cats      = exp_head_cat_totals[eh]
        eh_total  = sum(cats.values())
        eh_pct    = f"{(eh_total / grand_total * 100):.2f}%"

        # Exp head group header row
        for ci, val in enumerate([eh, "", round(eh_total, 2), eh_pct], 1):
            cell = ws_sum.cell(row=ri, column=ci, value=val)
            cell.fill = _hf; cell.font = _hft
            cell.alignment = Alignment(horizontal="center" if ci > 1 else "left")
        ri += 1

        # Category rows (sorted by amount desc)
        for cat, amt in sorted(cats.items(), key=lambda x: -x[1]):
            pct = f"{(amt / grand_total * 100):.2f}%"
            for ci, val in enumerate(["", cat, round(amt, 2), pct], 1):
                cell = ws_sum.cell(row=ri, column=ci, value=val)
                cell.fill = _cf
                cell.alignment = Alignment(horizontal="center" if ci > 1 else "left")
            ri += 1

        # Exp head TOTAL row
        for ci, val in enumerate([f"{eh} TOTAL", "", round(eh_total, 2), eh_pct], 1):
            cell = ws_sum.cell(row=ri, column=ci, value=val)
            cell.fill = _tf; cell.font = _tft
            cell.alignment = Alignment(horizontal="center" if ci > 1 else "left")
        ri += 1

    # Grand total row
    for ci, val in enumerate(["GRAND TOTAL", "", round(grand_total, 2), "100.00%"], 1):
        cell = ws_sum.cell(row=ri, column=ci, value=val)
        cell.fill = _gtf; cell.font = _gtft
        cell.alignment = Alignment(horizontal="center" if ci > 1 else "left")

    # ── FX REVALUATION sheet (only when non-INR outstanding items exist) ──────
    fx_outstanding = [] if base_ccy != "INR" else [
        b for b in bills
        if (b.get("currency_code") or "INR").strip() != "INR"
        and b.get("original_amount") is not None
        and (b.get("payment_status") or "unpaid") != "paid"
        and (b.get("status") or "").upper() not in ("VOID",)
    ]
    if fx_outstanding:
        ws_fx = wb.create_sheet("FX REVALUATION")
        fx_headers = [
            "VENDOR", "INVOICE NO", "CURRENCY", "ORIGINAL AMOUNT",
            "BOOKED RATE", "BOOKED INR", "CURRENT RATE", "CURRENT INR",
            "FX GAIN/LOSS", "DUE DATE", "STATUS",
        ]
        for ci, h in enumerate(fx_headers, 1):
            cell = ws_fx.cell(row=1, column=ci, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        total_booked = 0.0
        total_current = 0.0
        total_gl = 0.0

        for ri_fx, b in enumerate(fx_outstanding, 2):
            ccy = (b.get("currency_code") or "").strip()
            orig = float(b["original_amount"]) if b.get("original_amount") is not None else 0
            booked_rate = float(b["fx_rate"]) if b.get("fx_rate") is not None else None
            booked_inr = float(b["total_amount"]) if b.get("total_amount") is not None else 0
            current_rate = float(b["fx_rate_current"]) if b.get("fx_rate_current") is not None else booked_rate
            current_inr = float(b["inr_amount_current"]) if b.get("inr_amount_current") is not None else (
                orig * current_rate if current_rate else booked_inr
            )
            gl = float(b["fx_gain_loss"]) if b.get("fx_gain_loss") is not None else (current_inr - booked_inr)
            due = b.get("due_date")
            due_str = due.strftime("%d-%b-%Y") if due else ""

            total_booked += booked_inr
            total_current += current_inr
            total_gl += gl

            row_vals = [
                b.get("vendor_name") or "",
                str(b.get("stg_bill_id") or b.get("invoice_id") or ""),
                ccy,
                round(orig, 2),
                round(booked_rate, 6) if booked_rate else None,
                round(booked_inr, 2),
                round(current_rate, 6) if current_rate else None,
                round(current_inr, 2),
                round(gl, 2),
                due_str,
                b.get("status") or "",
            ]
            for ci, val in enumerate(row_vals, 1):
                cell = ws_fx.cell(row=ri_fx, column=ci, value=val)
                if ci == 9 and isinstance(val, (int, float)):
                    cell.font = Font(bold=True, color="059669" if val >= 0 else "DC2626")

        # Totals row
        tr = len(fx_outstanding) + 2
        for ci, val in enumerate(["TOTAL", "", "", "", "", round(total_booked, 2), "", round(total_current, 2), round(total_gl, 2), "", ""], 1):
            cell = ws_fx.cell(row=tr, column=ci, value=val)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F0F4FF")
            if ci == 9 and isinstance(val, (int, float)):
                cell.font = Font(bold=True, color="059669" if val >= 0 else "DC2626")

    # Auto-width columns
    all_sheets = [ws_raw, ws_sum]
    if fx_outstanding:
        all_sheets.append(ws_fx)
    for ws in all_sheets:
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=0)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"Expense_Report_{month_label}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

@router.get("/api/reports/expense-csv")
def expense_csv(
    company_id: str | None = Query(None),
    month: str | None = Query(None),   # e.g. "2026-04"
    status: str | None = Query(None),  # e.g. "APPROVED"
    report: str = Query("summary"),    # "summary" | "raw" | "both"
    brand_id: str | None = Query(None),
):
    """
    Generate vendor expense report in CSV format.
    - report=summary  → single CSV with EXP HEAD totals (default)
    - report=raw      → single CSV with one row per bill
    - report=both     → ZIP containing RAW + SUMMARY CSVs
    """
    import psycopg2.extras
    import io
    import csv as _csv
    import zipfile
    from datetime import date
    from fastapi.responses import StreamingResponse, Response

    tenant = require_tenant()  # tenant from JWT; client company_id ignored

    # ── Date range (month=None or "all" → no date filter) ───────────────────
    use_date_filter = bool(month and month.lower() != "all")
    if use_date_filter:
        year, mon = int(month.split("-")[0]), int(month.split("-")[1])
        first_day = date(year, mon, 1)
        next_first = date(year, mon + 1, 1) if mon < 12 else date(year + 1, 1, 1)
        month_label = first_day.strftime("%b-%y").upper()
    else:
        month_label = "ALL"

    conn = get_conn()
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            # Fetch bills
            sql = """
                SELECT
                    fi.invoice_id,
                    fi.stg_bill_id,
                    fi.invoice_date,
                    fi.invoice_number,
                    fi.total_amount,
                    fi.category,
                    fi.status,
                    fi.amount_paid,
                    fi.outstanding_amount,
                    ds.name AS vendor_name,
                    ds.gstn AS vendor_gstn
                FROM fact_invoices fi
                LEFT JOIN dim_suppliers ds ON ds.supplier_id = fi.supplier_id
                WHERE fi.tenant_id = %s
            """
            params: list = [tenant]
            if use_date_filter:
                sql += " AND fi.invoice_date >= %s AND fi.invoice_date < %s"
                params.extend([first_day, next_first])
            if status:
                sql += " AND UPPER(fi.status::text) = %s"
                params.append(status.upper())
            if brand_id:
                sql += " AND fi.brand_id = %s::uuid"
                params.append(brand_id)
            sql += " ORDER BY fi.invoice_date, ds.name"
            cur.execute(sql, params)
            bills = cur.fetchall()

            # COA rules (highest priority first)
            cur.execute(
                """
                SELECT exp_head, vendor_name, category, priority
                FROM coa_rules
                WHERE tenant_id = %s
                ORDER BY priority DESC
                """,
                (tenant,),
            )
            rules = cur.fetchall()

            # dim_categories exp_head assignments (4th tier fallback)
            cur.execute(
                "SELECT name, exp_head FROM dim_categories WHERE tenant_id = %s AND exp_head IS NOT NULL",
                (tenant,),
            )
            cat_exp_heads_csv = {r["name"].lower(): r["exp_head"] for r in cur.fetchall()}

            # Company name — use brand name if brand_id filter is active, else tenant name
            if brand_id:
                cur.execute("SELECT name FROM brands WHERE id = %s", (brand_id,))
                b_row = cur.fetchone()
                company_name = (b_row["name"] if b_row else "") or "COMPANY"
            else:
                cur.execute("SELECT company_name FROM dim_tenants WHERE tenant_id = %s", (tenant,))
                t_row = cur.fetchone()
                company_name = (t_row["company_name"] if t_row else "") or "COMPANY"
    finally:
        put_conn(conn)

    # ── COA resolution ───────────────────────────────────────────────────────
    def resolve_exp_head(vendor: str | None, category: str | None) -> str:
        v = (vendor or "").lower()
        c = (category or "").lower()
        for r in rules:
            rv = (r["vendor_name"] or "").lower()
            rc = (r["category"] or "").lower()
            if rv and rc and rv in v and rc == c:
                return r["exp_head"]
        for r in rules:
            rv = (r["vendor_name"] or "").lower()
            if rv and not r["category"] and rv in v:
                return r["exp_head"]
        for r in rules:
            rc = (r["category"] or "").lower()
            if not r["vendor_name"] and rc and rc == c:
                return r["exp_head"]
        if c in cat_exp_heads_csv:
            return cat_exp_heads_csv[c]
        return "OTHER"

    # ── Build row data ───────────────────────────────────────────────────────
    raw_rows: list[list] = []
    exp_head_cat_totals: dict[str, dict[str, float]] = {}

    for bill in bills:
        vendor   = (bill["vendor_name"] or "").strip()
        category = (bill["category"] or "General").strip()
        exp_head = resolve_exp_head(vendor, category)
        total    = float(bill["total_amount"] or 0)
        status_val = (str(bill["status"]) if bill["status"] else "").upper()
        opex = "Y" if status_val in ("PAID", "PARTIAL") else ""
        inv_date  = bill["invoice_date"]
        svc_month = inv_date.strftime("%b-%y").upper() if inv_date else month_label
        inv_no    = str(bill["stg_bill_id"] or bill["invoice_id"])

        raw_rows.append([
            company_name, vendor, inv_no, svc_month,
            vendor, category, exp_head,
            round(total, 2), opex, round(total, 2),
        ])
        eh_cats = exp_head_cat_totals.setdefault(exp_head, {})
        eh_cats[category] = eh_cats.get(category, 0.0) + total

    grand_total = sum(sum(c.values()) for c in exp_head_cat_totals.values())

    RAW_HEADERS = [
        "COMPANY", "VENDOR NAME", "INVOICE NO", "SERVICE MONTH",
        "SUB HEAD", "SERVICE (CATEGORY)", "EXP HEAD",
        "TOTAL", "OPEX", "MONTHLY EXPENSE",
    ]
    SUM_HEADERS = ["EXP HEAD", "CATEGORY", "AMOUNT", "EXP RATIO IN %"]

    # ── Helper: build CSV buffer ─────────────────────────────────────────────
    def _raw_csv_buf() -> io.StringIO:
        buf = io.StringIO()
        w = _csv.writer(buf)
        w.writerow(RAW_HEADERS)
        w.writerows(raw_rows)
        buf.seek(0)
        return buf

    def _summary_csv_buf() -> io.StringIO:
        buf = io.StringIO()
        w = _csv.writer(buf)
        w.writerow(SUM_HEADERS)
        sorted_heads = sorted(exp_head_cat_totals.keys(), key=lambda h: (h == "OTHER", h))
        for eh in sorted_heads:
            cats     = exp_head_cat_totals[eh]
            eh_total = sum(cats.values())
            for cat, amt in sorted(cats.items(), key=lambda x: -x[1]):
                pct = (amt / grand_total * 100) if grand_total else 0.0
                w.writerow([eh, cat, round(amt, 2), round(pct, 2)])
            eh_pct = (eh_total / grand_total * 100) if grand_total else 0.0
            w.writerow([f"{eh} TOTAL", "", round(eh_total, 2), round(eh_pct, 2)])
        w.writerow(["GRAND TOTAL", "", round(grand_total, 2), "100.00"])
        buf.seek(0)
        return buf

    # ── Return based on report type ──────────────────────────────────────────
    report = report.lower()

    if report == "raw":
        content = _raw_csv_buf().getvalue().encode("utf-8-sig")
        fname = f"Expense_Report_RAW_{month_label}.csv"
        return Response(
            content=content,
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{fname}"'},
        )

    if report == "both":
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(
                f"Expense_Report_RAW_{month_label}.csv",
                _raw_csv_buf().getvalue(),
            )
            zf.writestr(
                f"Expense_Report_SUMMARY_{month_label}.csv",
                _summary_csv_buf().getvalue(),
            )
        zip_buf.seek(0)
        fname = f"Expense_Report_{month_label}.zip"
        return StreamingResponse(
            zip_buf,
            media_type="application/zip",
            headers={"Content-Disposition": f'attachment; filename="{fname}"'},
        )

    # Default: summary
    content = _summary_csv_buf().getvalue().encode("utf-8-sig")
    fname = f"Expense_Report_SUMMARY_{month_label}.csv"
    return Response(
        content=content,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
